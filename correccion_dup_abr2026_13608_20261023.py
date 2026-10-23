#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Corrección de ingresos Bancolombia DUPLICADOS de abril 2026 de Julian Sanchez (13608), aprobada por Julián
el 23-oct-2026 ("primero valida que todos sigan ahí, segundo cárgalo").

Auditoría (Downloads/Auditoria_mayoristas_2026/13608 - Julian Sanchez): en abril el CSV diario del banco
traía algunas transacciones dos veces (en línea + asentada) y el cargue las tomó como distintas →
9 ingresos por COP 29.799.999 que NO están en el extracto oficial (cuenta …4286), 1 a 11-abr.

Decisión: NO se borran las filas de abril (los saldos históricos no se reprocesan). Se aplica UNA
corrección con fecha de HOY:
  13608: Ingreso / Ingreso_extra, Monto -29.799.999, Orden correccion_dup_bcol_abr2026_13608, Fecha 23-oct.
También se agrega a ingresos_extra.xlsx (OneDrive) con el mismo Orden (no ExtraN), para que re-subir ese
archivo no la pierda ni la duplique (dedup de Ingreso_extra por Orden+Motivo).

  sin argumentos  -> dry-run (0 escrituras)
  --escribir REV  -> escribe a Dropbox (REV = rev del dry-run), ingresos_extra.xlsx y la copia de OneDrive
"""
import sys, os, io, shutil, warnings
warnings.filterwarnings("ignore")
import pandas as pd

ESCRIBIR = "--escribir" in sys.argv
REV_ESPERADA = sys.argv[sys.argv.index("--escribir") + 1] if ESCRIBIR else None
CAS = "13608"
ORDEN_ORIG = "correccion_dup_bcol_abr2026_13608"
MONTO = 29799999.0
FECHA = pd.Timestamp("2026-10-23")
NOMBRE = ("Corrección conciliación abril 2026: 9 ingresos Bancolombia registrados dos veces "
          "(1 a 11-abr) que no aparecen en el extracto oficial de la cuenta terminada en 4286")
AUD = "/Users/julianlopez/Downloads/Auditoria_mayoristas_2026/13608 - Julian Sanchez/Ingresos_abril_2026_13608.xlsx"
EXTRA = "/Users/julianlopez/Library/CloudStorage/OneDrive-Personal/Historico Carga/Conciliacion/Mayoristas/ingresos_extra.xlsx"
ONEDRIVE_DIR = "/Users/julianlopez/Library/CloudStorage/OneDrive-Personal/Historico Carga/Conciliacion/Mayoristas"
ONEDRIVE_ANT = f"{ONEDRIVE_DIR}/Antiguos"


def saldo(d):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    return float(pd.to_numeric(t["Monto"], errors="coerce").iloc[-1]) if len(t) else float("nan")


def usuario_de_totales(d):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    u = t["Usuario"].astype(str).str.strip()
    u = u[~u.str.lower().isin({"", "nan", "none"})]
    return u.mode().iloc[0] if len(u) else ""


def cas_de_totales(d, fb):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    v = t["Casillero"].dropna()
    return v.iloc[-1] if len(v) else fb


def main():
    import harness
    mod = harness.cargar_app()
    SEP = "=" * 94
    def banner(t): print(f"\n{SEP}\n{t}\n{SEP}")
    ok = True
    def chk(n, c, det=""):
        nonlocal ok
        print(f"  {'✔' if c else '🚨'} {n:<66} {det}")
        ok = ok and bool(c)
    print(f"MODO: {'🔴 ESCRITURA REAL' if ESCRIBIR else '🟢 DRY-RUN (0 escrituras)'}")
    for fn in ("procesar_egresos", "procesar_amex", "procesar_rakuten", "procesar_robinhood",
               "procesar_capital", "procesar_usbank", "procesar_intuit", "procesar_discover",
               "agregar_incentivo_amex"):
        def _boom(*a, _n=fn, **k):
            raise AssertionError(f"{_n} fue llamada — este script solo aplica una corrección")
        setattr(mod, fn, _boom)

    banner("2) HISTÓRICO VIVO FRESCO")
    cfg = mod.st.secrets["dropbox"]
    md = mod.dbx.files_get_metadata(cfg["remote_path"])
    _, res = mod.dbx.files_download(cfg["remote_path"])
    vivo = pd.read_excel(io.BytesIO(res.content), sheet_name=None)
    HOJA = next(h for h in vivo if h.split(" - ")[0].strip() == CAS)
    s0 = saldo(vivo[HOJA])
    print(f"  rev={md.rev}  modificado={md.server_modified}  size={md.size:,}")
    print(f"  '{HOJA}': saldo {s0:,.2f}")
    if ESCRIBIR and md.rev != REV_ESPERADA:
        raise SystemExit(f"⛔ ABORTA: el histórico se movió (esperaba {REV_ESPERADA}). Rehacer dry-run.")

    banner("3) LOS 9 DUPLICADOS DE LA AUDITORÍA (se verifican, NO se borran)")
    dup = pd.read_excel(AUD, sheet_name="Duplicados")
    ids = set(dup["ID_INGRESO"].astype(str).str.strip())
    o0 = vivo[HOJA]["Orden"].astype(str).str.strip()
    idv = vivo[HOJA]["ID_INGRESO"].astype(str).str.strip()
    r = vivo[HOJA][idv.isin(ids)]
    print(r[["Fecha", "Tipo", "Monto", "ID_INGRESO", "Nombre del producto"]].to_string())
    chk("la auditoría trae 9 duplicados por 29.799.999",
        len(dup) == 9 and abs(float(dup["Monto"].sum()) - MONTO) < 0.01)
    chk("los 9 siguen en 13608, una vez cada uno, como Ingreso, Σ 29.799.999",
        len(r) == 9 and r["ID_INGRESO"].astype(str).str.strip().nunique() == 9
        and set(r["Tipo"].astype(str).str.strip()) == {"Ingreso"}
        and abs(float(pd.to_numeric(r["Monto"]).sum()) - MONTO) < 0.01)
    chk("todos son de abril 2026", set(pd.to_datetime(r["Fecha"]).dt.strftime("%Y-%m")) == {"2026-04"})
    chk("13608 todavía no tiene la corrección", not (o0 == ORDEN_ORIG).any())
    ex = pd.read_excel(EXTRA, sheet_name=None)
    HEX = next(h for h in ex if h.split(" - ")[0].strip() == CAS)
    chk("ingresos_extra.xlsx todavía no tiene la corrección",
        not (ex[HEX]["Orden"].astype(str).str.strip() == ORDEN_ORIG).any(), f"hoja '{HEX}' {len(ex[HEX])} filas")
    if not ok:
        raise SystemExit("⛔ ABORTA.")

    banner("4) APLICAR")
    hoy = pd.Timestamp.today().strftime("%Y-%m-%d")
    historico = {k: mod.asegurar_columnas_historico(v.copy()) for k, v in vivo.items()}
    d = historico[HOJA]
    u_orig = usuario_de_totales(vivo[HOJA]) or "Julian Sanchez"
    fila = {c: "" for c in d.columns}
    fila.update({"Fecha": FECHA, "Tipo": "Ingreso", "Monto": -MONTO, "Orden": ORDEN_ORIG,
                 "Motivo": "Ingreso_extra", "TRM": "", "Usuario": u_orig, "Casillero": int(CAS),
                 "Nombre del producto": NOMBRE, "Fecha de Carga": hoy})
    d = pd.concat([d, pd.DataFrame([fila])], ignore_index=True)
    historico[HOJA] = mod.recalcular_totales_diarios(d, usuario=u_orig, cas=cas_de_totales(vivo[HOJA], CAS))
    print(historico[HOJA][historico[HOJA]["Orden"].astype(str) == ORDEN_ORIG][["Fecha","Tipo","Monto","Orden","Motivo","Usuario"]].to_string())
    def tots(df):
        t = df[df["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"].copy()
        return t.assign(F=pd.to_datetime(t["Fecha"]).dt.normalize()).groupby("F")["Monto"].last().astype(float)
    ta, tb = tots(vivo[HOJA]), tots(historico[HOJA])
    print(pd.DataFrame({"antes": ta, "después": tb}).tail(6).to_string(float_format=lambda v: f"{v:,.2f}"))

    banner("5) INVARIANTES")
    s1 = saldo(historico[HOJA])
    chk("Julián Sánchez baja exactamente 29.799.999", abs((s1 - s0) + MONTO) < 0.01, f"{s0:,.2f} → {s1:,.2f}")
    def movs(df):
        return int((df["Tipo"].astype(str).str.strip().str.upper() != "TOTAL").sum())
    chk("Julián Sánchez: +1 fila de movimiento (no se borra nada)", movs(historico[HOJA]) == movs(vivo[HOJA]) + 1)
    comun = ta.index[ta.index < FECHA.normalize()]
    chk("todos los saldos diarios ANTERIORES al 23-oct quedan idénticos",
        (tb.reindex(comun) - ta.reindex(comun)).abs().max() < 0.01, f"{len(comun)} días")
    vac = {"", "nan", "none", "nat"}
    for h in (HOJA,):
        a = {y for y in vivo[h]["Orden"].astype(str).str.strip() if y.lower() not in vac}
        b = {y for y in historico[h]["Orden"].astype(str).str.strip() if y.lower() not in vac}
        chk(f"{h}: 0 Orden perdidos", not (a - b), f"{len(a - b)}")
    for hoja in vivo:
        if hoja == HOJA:
            continue
        A, B = vivo[hoja], historico[hoja]
        igual = len(A) == len(B)
        if igual and "Monto" in A.columns:
            igual = abs(pd.to_numeric(A["Monto"], errors="coerce").fillna(0).sum() -
                        pd.to_numeric(B["Monto"], errors="coerce").fillna(0).sum()) < 0.01
        chk(f"verbatim: {hoja}", igual, f"{len(A)} filas")
    if not ok:
        raise SystemExit("⛔ ABORTA: falló alguna invariante.")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for hh, dfh in historico.items():
            w.book.create_sheet(hh[:31])
            dfh.to_excel(w, sheet_name=hh[:31], index=False)
    buf.seek(0); data_bytes = buf.read()
    harness.clear_msgs()
    try:
        mod.guard_frescura_historico(historico)
        print("  ✅ GUARD A PASA (0 pérdidas)")
    except harness._Stop:
        for n, t in harness.MENSAJES:
            print(f"  [{n}] {t[:400]}")
        raise SystemExit("⛔ GUARD A BLOQUEÓ")

    if not ESCRIBIR:
        banner("DRY-RUN TERMINADO — 0 ESCRITURAS")
        print(f"  Para escribir: python3 correccion_dup_abr2026_13608_20261023.py --escribir {md.rev}")
        return

    banner("7) SUBIDA A DROPBOX (capa C hace el respaldo)")
    md_pre = mod.dbx.files_get_metadata(cfg["remote_path"])
    if md_pre.rev != REV_ESPERADA:
        raise SystemExit(f"⛔ ABORTA SIN ESCRIBIR: el histórico se movió (rev {md_pre.rev}).")
    harness.clear_msgs()
    mod.upload_to_dropbox(data_bytes)
    backup = None
    for n, t in harness.MENSAJES:
        print(f"  [{n}] {t}")
        if "Respaldo previo creado" in t and "`" in t:
            backup = t.split("`")[1]
    harness.clear_msgs()

    banner("8) VALIDACIÓN POST-ESCRITURA")
    md2 = mod.dbx.files_get_metadata(cfg["remote_path"])
    print(f"  rev NUEVA = {md2.rev}  size = {md2.size:,}")
    print(f"  backup    = {backup}")
    _, r2 = mod.dbx.files_download(cfg["remote_path"])
    contenido = r2.content
    rel = pd.read_excel(io.BytesIO(contenido), sheet_name=None)
    ok2 = True
    def chk2(n, c, det=""):
        nonlocal ok2
        print(f"  {'✔' if c else '🚨'} {n:<64} {det}")
        ok2 = ok2 and bool(c)
    qa = rel[HOJA]["Orden"].astype(str).str.strip() == ORDEN_ORIG
    chk2("13608: corrección de -29.799.999 presente", int(qa.sum()) == 1 and abs(float(pd.to_numeric(rel[HOJA].loc[qa, "Monto"]).sum()) + MONTO) < 0.01)
    chk2("13608: saldo", abs(saldo(rel[HOJA]) - s1) < 0.01, f"COP {saldo(rel[HOJA]):,.2f}")
    for hoja in vivo:
        if hoja == HOJA:
            continue
        A, B = vivo[hoja], rel[hoja]
        igual = len(A) == len(B)
        if igual and "Monto" in A.columns:
            igual = abs(pd.to_numeric(A["Monto"], errors="coerce").fillna(0).sum() -
                        pd.to_numeric(B["Monto"], errors="coerce").fillna(0).sum()) < 0.01
        chk2(f"verbatim: {hoja}", igual, f"{len(A)} filas")

    banner("9) ingresos_extra.xlsx (OneDrive) — misma fila, mismo Orden")
    import openpyxl
    bk = EXTRA.replace(".xlsx", f"_backup_{pd.Timestamp.today():%Y%m%d_%H%M%S}_pre_correccion_13608.xlsx")
    shutil.copy2(EXTRA, bk)
    chk2("respaldo de ingresos_extra.xlsx", os.path.getsize(bk) == os.path.getsize(EXTRA), bk)
    wb = openpyxl.load_workbook(EXTRA); ws = wb[HEX]
    cab = [c.value for c in ws[1]]
    val = {"Tipo": "Ingreso", "Fecha": FECHA.to_pydatetime(), "Orden": ORDEN_ORIG, "Monto": -MONTO,
           "Usuario": u_orig, "Casillero": int(CAS), "Motivo": "Ingreso_extra", "Nombre del producto": NOMBRE}
    ws.append([val.get(c) for c in cab]); wb.save(EXTRA)
    ex2 = pd.read_excel(EXTRA, sheet_name=None)
    chk2("ingresos_extra.xlsx: la corrección quedó una vez", int((ex2[HEX]["Orden"].astype(str) == ORDEN_ORIG).sum()) == 1)
    chk2("ingresos_extra.xlsx: las demás hojas intactas", all(ex2[k].equals(ex[k]) for k in ex if k != HEX))

    banner("10) COPIA A ONEDRIVE")
    hoy = pd.Timestamp.today().strftime("%Y%m%d")
    destino = f"{ONEDRIVE_DIR}/{hoy}_Historico_mayoristas.xlsx"
    os.makedirs(ONEDRIVE_ANT, exist_ok=True)
    # ⚠️ COPIAR, no mover (el 18-sep OneDrive perdió un archivo movido): copiar, verificar
    # tamaño y solo entonces borrar/sobrescribir.
    for f in sorted(os.listdir(ONEDRIVE_DIR)):
        if f.endswith("_Historico_mayoristas.xlsx"):
            src_f = f"{ONEDRIVE_DIR}/{f}"
            dst_f = f"{ONEDRIVE_ANT}/{f.replace('.xlsx','')}_PRE_correccion_13608.xlsx"
            if os.path.exists(dst_f):
                raise SystemExit(f"⛔ ya existe {dst_f}; no se pisa. Revisar a mano.")
            shutil.copy2(src_f, dst_f)
            if not (os.path.exists(dst_f) and os.path.getsize(dst_f) == os.path.getsize(src_f)):
                raise SystemExit(f"⛔ la copia a Antiguos no quedó bien; NO se reemplaza {f}.")
            if f != os.path.basename(destino):
                os.remove(src_f)
            print(f"  archivada: {f} → Antiguos/{os.path.basename(dst_f)}")
    with open(destino, "wb") as fh:
        fh.write(contenido)
    print(f"  escrita:   {destino}  ({os.path.getsize(destino):,} bytes)")
    chk2("la copia de OneDrive pesa lo mismo que el vivo",
         os.path.getsize(destino) == md2.size, f"{os.path.getsize(destino):,} vs {md2.size:,}")

    print(f"\n  {'✅ CORRECCIÓN APLICADA A 13608 Y VERIFICADA' if ok2 else '🚨 REVISAR'}")
    print(f"  rev nueva: {md2.rev}\n  rollback:  {backup}")
    print("\n  ↩️  Para revertir: borrar la fila correccion_dup_bcol_abr2026_13608 (histórico e ingresos_extra.xlsx) y recalcular, o restaurar los backups.")


if __name__ == "__main__":
    sys.path.insert(0, os.environ.get("HARNESS_DIR", "."))
    main()
