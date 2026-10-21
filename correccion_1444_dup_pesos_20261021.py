#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
1444 · María Moises — dos correcciones aprobadas por Julián el 21-oct-2026 ("listo, hazlo"):

A) PESOS FINALES de 186 envíos de septiembre (archivo 20261019_Comparacion_pesos_1444_resumida.xlsx):
   se REEMPLAZA el peso en el propio envío (Peso_lb) y la app recalcula el Monto con su tarifa
   (max(lb,1)×6+5 USD × TRM de la fila). La regla PESO_FINAL_1444 de la app lo mantiene en cada corrida.

B) DUPLICADOS de abril 2026: 141 ingresos Bancolombia (cuenta …2930) por COP 28.921.400.

   NO se borran las filas de abril: un solo cobro total con fecha de HOY:
   1444: Ingreso / Ingreso_extra, Monto -28.921.400, Orden correccion_dup_bcol_abr2026_1444, Fecha 21-oct.
También se agrega a ingresos_extra.xlsx (OneDrive) con el mismo Orden (no ExtraN), para que re-subir ese
archivo no la pierda ni la duplique (dedup de Ingreso_extra por Orden+Motivo).

  sin argumentos  -> dry-run (0 escrituras)
  --escribir REV  -> escribe a Dropbox (REV = rev del dry-run), ingresos_extra.xlsx y la copia de OneDrive
"""
import sys, os, io, shutil, warnings
warnings.filterwarnings("ignore")
import pandas as pd

ESCRIBIR = "--escribir" in sys.argv
REV_ESPERADA = sys.argv[sys.argv.index("--escribir") + 1] if ESCRIBIR else None
CAS = "1444"
ORDEN_ORIG = "correccion_dup_bcol_abr2026_1444"
MONTO = 28921400.0
FECHA = pd.Timestamp("2026-10-21")
NOMBRE = ("Corrección conciliación abril 2026: 141 ingresos Bancolombia registrados dos veces "
          "(6, 8, 9, 10, 11 y 13-abr) que no aparecen en el extracto oficial de la cuenta terminada en 2930")
AUD = "/Users/julianlopez/Downloads/Auditoria_mayoristas_2026/1444 - Maria Moises/Ingresos_abril_2026_1444.xlsx"
COMP = "/private/tmp/claude-501/-Users-julianlopez/0c0e43d6-95b4-481a-a558-31041bfb4896/scratchpad/pesos/comp.xlsx"  # copia de ~/Downloads/20261019_Comparacion_pesos_1444_resumida (1).xlsx
EXTRA = "/Users/julianlopez/Library/CloudStorage/OneDrive-Personal/Historico Carga/Conciliacion/Mayoristas/ingresos_extra.xlsx"
ONEDRIVE_DIR = "/Users/julianlopez/Library/CloudStorage/OneDrive-Personal/Historico Carga/Conciliacion/Mayoristas"
ONEDRIVE_ANT = f"{ONEDRIVE_DIR}/Antiguos"


def saldo(d):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    return float(pd.to_numeric(t["Monto"], errors="coerce").iloc[-1]) if len(t) else float("nan")


def usuario_de_totales(d):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    u = t["Usuario"].astype(str).str.strip()
    u = u[~u.str.lower().isin({"", "nan", "none"})]
    return u.mode().iloc[0] if len(u) else ""


def cas_de_totales(d, fb):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    v = t["Casillero"].dropna()
    return v.iloc[-1] if len(v) else fb


def main():
    import harness
    mod = harness.cargar_app()
    SEP = "=" * 94
    def banner(t): print(f"\n{SEP}\n{t}\n{SEP}")
    ok = True
    def chk(n, c, det=""):
        nonlocal ok
        print(f"  {'✔' if c else '🚨'} {n:<70} {det}")
        ok = ok and bool(c)
    print(f"MODO: {'🔴 ESCRITURA REAL' if ESCRIBIR else '🟢 DRY-RUN (0 escrituras)'}")
    for fn in ("procesar_egresos", "procesar_amex", "procesar_rakuten", "procesar_robinhood",
               "procesar_capital", "procesar_usbank", "procesar_intuit", "procesar_discover",
               "agregar_incentivo_amex"):
        def _boom(*a, _n=fn, **k):
            raise AssertionError(f"{_n} fue llamada — este script solo aplica una corrección")
        setattr(mod, fn, _boom)

    banner("1) LA APP TRAE LA REGLA PESO_FINAL_1444")
    PF = mod.PESO_FINAL_1444
    comp = pd.read_excel(COMP)
    comp = comp[pd.to_numeric(comp["Envío"], errors="coerce").notna()].copy()
    comp["O"] = "Envio " + comp["Envío"].astype(int).astype(str)
    esp = comp[comp["Peso final (lb)"].notna() & (comp["Diferencia COP"].fillna(0) != 0)]
    chk("PESO_FINAL_1444 = los 186 envíos del archivo con su peso final",
        len(PF) == 186 and dict(zip(esp.O, esp["Peso final (lb)"].astype(float))) == {k: float(v) for k, v in PF.items()})

    banner("2) HISTÓRICO VIVO FRESCO")
    cfg = mod.st.secrets["dropbox"]
    md = mod.dbx.files_get_metadata(cfg["remote_path"])
    _, res = mod.dbx.files_download(cfg["remote_path"])
    vivo = pd.read_excel(io.BytesIO(res.content), sheet_name=None)
    HOJA = next(h for h in vivo if h.split(" - ")[0].strip() == CAS and h.strip().endswith("Moises"))
    s0 = saldo(vivo[HOJA])
    print(f"  rev={md.rev}  modificado={md.server_modified}  size={md.size:,}")
    print(f"  '{HOJA}': saldo {s0:,.2f}")
    if ESCRIBIR and md.rev != REV_ESPERADA:
        raise SystemExit(f"⛔ ABORTA: el histórico se movió (esperaba {REV_ESPERADA}). Rehacer dry-run.")

    banner("3) VERIFICACIONES PREVIAS")
    dup = pd.read_excel(AUD, sheet_name="Duplicados")
    ids = set(dup["ID_INGRESO"].astype(str).str.strip())
    o0 = vivo[HOJA]["Orden"].astype(str).str.strip()
    idv = vivo[HOJA]["ID_INGRESO"].astype(str).str.strip()
    r = vivo[HOJA][idv.isin(ids)]
    chk("auditoría: 141 duplicados por 28.921.400",
        len(dup) == 141 and abs(float(dup["Monto"].sum()) - MONTO) < 0.01)
    chk("los 141 siguen en 1444, una vez cada uno, como Ingreso, Σ 28.921.400",
        len(r) == 141 and r["ID_INGRESO"].astype(str).str.strip().nunique() == 141
        and set(r["Tipo"].astype(str).str.strip()) == {"Ingreso"}
        and abs(float(pd.to_numeric(r["Monto"]).sum()) - MONTO) < 0.01)
    chk("todos son de abril 2026", set(pd.to_datetime(r["Fecha"]).dt.strftime("%Y-%m")) == {"2026-04"})
    chk("1444 todavía no tiene la corrección de duplicados", not (o0 == ORDEN_ORIG).any())
    ex = pd.read_excel(EXTRA, sheet_name=None)
    HEX = next(h for h in ex if h.split(" - ")[0].strip() == CAS)
    chk("ingresos_extra.xlsx todavía no la tiene",
        not (ex[HEX]["Orden"].astype(str).str.strip() == ORDEN_ORIG).any(), f"hoja '{HEX}' {len(ex[HEX])} filas")
    env = vivo[HOJA][o0.isin(PF)]
    chk("los 186 envíos están en 1444, una vez cada uno, como Egreso/Envio",
        len(env) == 186 and env["Orden"].nunique() == 186
        and set(env["Motivo"].astype(str).str.strip()) == {"Envio"})
    m = esp.merge(env[["Orden", "Monto", "TRM"]], left_on="O", right_on="Orden")
    chk("Monto actual = USD cobrado × TRM del archivo (±1 COP)",
        len(m) == 186 and ((pd.to_numeric(m["Monto"]) - m["USD cobrado"] * m["TRM_x"]).abs() <= 1).all())
    chk("TRM del histórico = TRM del archivo", ((pd.to_numeric(m["TRM_y"]) - m["TRM_x"]).abs() < 0.01).all())
    if not ok:
        raise SystemExit("⛔ ABORTA.")

    banner("4) APLICAR")
    hoy = pd.Timestamp.today().strftime("%Y-%m-%d")
    historico = {k: mod.asegurar_columnas_historico(v.copy()) for k, v in vivo.items()}
    d = historico[HOJA]
    antes = d.copy()
    # A) pesos finales: la MISMA función que corre la app (con la regla PESO_FINAL_1444)
    d, tp = mod.aplicar_tarifa_envio_por_peso(d, CAS, conf=mod.TARIFA_ENVIO_1444, desde=mod.TARIFA_ENVIO_1444_DESDE)
    print(f"  tarifa por peso: {tp['recalculadas']} recalculadas · {tp['subieron']} suben · {tp['bajaron']} bajan · COP {tp['cop']:+,.0f}")
    ma, mb = pd.to_numeric(antes["Monto"], errors="coerce"), pd.to_numeric(d["Monto"], errors="coerce")
    camb = (ma - mb).abs() > 0.001
    chk("cambian exactamente los Monto de los 186 envíos (y ninguna otra fila)",
        set(d.loc[camb, "Orden"].astype(str).str.strip()) == set(PF) and int(camb.sum()) == 186)
    dl = (mb - ma)[camb]
    chk("182 suben y 4 bajan", int((dl > 0).sum()) == 182 and int((dl < 0).sum()) == 4)
    dd = esp.set_index("O")["Diferencia COP"]
    por = pd.Series(dl.values, index=d.loc[camb, "Orden"].astype(str).str.strip().values)
    chk("cada envío cambia lo mismo que dice el archivo (±1 COP por redondeo)",
        (por - dd.reindex(por.index)).abs().max() <= 1, f"máx dif {float((por - dd.reindex(por.index)).abs().max()):.2f}")
    DELTA_PESOS = float(dl.sum())
    print(f"  Δ pesos = {DELTA_PESOS:,.2f}  (archivo: {float(dd.sum()):,.2f})")
    chk("Peso_lb de los 186 = peso final", all(abs(float(d.loc[d["Orden"].astype(str).str.strip() == o, "Peso_lb"].iloc[0]) - float(p)) < 1e-9 for o, p in PF.items()))
    # B) cobro total de los duplicados
    u_orig = usuario_de_totales(vivo[HOJA]) or "Maria Moises"
    fila = {c: "" for c in d.columns}
    fila.update({"Fecha": FECHA, "Tipo": "Ingreso", "Monto": -MONTO, "Orden": ORDEN_ORIG,
                 "Motivo": "Ingreso_extra", "TRM": "", "Usuario": u_orig, "Casillero": int(CAS),
                 "Nombre del producto": NOMBRE, "Fecha de Carga": hoy})
    d = pd.concat([d, pd.DataFrame([fila])], ignore_index=True)
    historico[HOJA] = mod.recalcular_totales_diarios(d, usuario=u_orig, cas=cas_de_totales(vivo[HOJA], CAS))
    def tots(df):
        t = df[df["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"].copy()
        return t.assign(F=pd.to_datetime(t["Fecha"]).dt.normalize()).groupby("F")["Monto"].last().astype(float)
    ta, tb = tots(vivo[HOJA]), tots(historico[HOJA])
    print(pd.DataFrame({"antes": ta, "después": tb}).tail(8).to_string(float_format=lambda v: f"{v:,.2f}"))

    banner("5) INVARIANTES")
    s1 = saldo(historico[HOJA])
    chk("saldo baja exactamente duplicados + pesos", abs((s0 - s1) - (MONTO + DELTA_PESOS)) < 0.01,
        f"{s0:,.2f} → {s1:,.2f}")
    chk("queda con saldo POSITIVO", s1 > 0)
    chk("todos los días desde el 16-oct siguen positivos", (tb[tb.index >= "2026-10-16"] > 0).all(),
        f"mín {tb[tb.index >= '2026-10-16'].min():,.2f}")
    def movs(df):
        return int((df["Tipo"].astype(str).str.strip().str.upper() != "TOTAL").sum())
    chk("+1 fila de movimiento (no se borra nada)", movs(historico[HOJA]) == movs(vivo[HOJA]) + 1)
    comun = ta.index[ta.index < pd.Timestamp("2026-09-04")]
    chk("saldos diarios ANTERIORES al 4-sep idénticos", (tb.reindex(comun) - ta.reindex(comun)).abs().max() < 0.01, f"{len(comun)} días")
    def comis(df):
        c = df[df["Orden"].astype(str).str.contains("Comision", case=False, na=False)]
        return sorted(zip(c["Orden"].astype(str), pd.to_numeric(c["Monto"]).round(2)))
    chk("filas de comisión idénticas", comis(vivo[HOJA]) == comis(historico[HOJA]), f"{len(comis(vivo[HOJA]))}")
    vac = {"", "nan", "none", "nat"}
    a = {y for y in vivo[HOJA]["Orden"].astype(str).str.strip() if y.lower() not in vac}
    b = {y for y in historico[HOJA]["Orden"].astype(str).str.strip() if y.lower() not in vac}
    chk(f"{HOJA}: 0 Orden perdidos", not (a - b), f"{len(a - b)}")
    for hoja in vivo:
        if hoja == HOJA:
            continue
        A, B = vivo[hoja], historico[hoja]
        igual = len(A) == len(B)
        if igual and "Monto" in A.columns:
            igual = abs(pd.to_numeric(A["Monto"], errors="coerce").fillna(0).sum() -
                        pd.to_numeric(B["Monto"], errors="coerce").fillna(0).sum()) < 0.01
        chk(f"verbatim: {hoja}", igual, f"{len(A)} filas")
    # punto fijo: correr la regla otra vez no cambia nada
    d3, tp3 = mod.aplicar_tarifa_envio_por_peso(historico[HOJA], CAS, conf=mod.TARIFA_ENVIO_1444, desde=mod.TARIFA_ENVIO_1444_DESDE)
    chk("idempotente: una 2ª pasada de la tarifa no cambia nada", tp3["recalculadas"] == 0)
    # simulación del próximo cargue: una fila del archivo de envíos con el peso viejo
    sim = historico[HOJA].copy()
    i0 = sim.index[sim["Orden"].astype(str).str.strip() == "Envio 105926"][0]
    sim.loc[i0, "Peso_lb"] = 1.0; sim.loc[i0, "Monto"] = 36497.0
    sim, _ = mod.aplicar_tarifa_envio_por_peso(sim, CAS, conf=mod.TARIFA_ENVIO_1444, desde=mod.TARIFA_ENVIO_1444_DESDE)
    chk("si el archivo de envíos trae el peso viejo, la regla lo vuelve al final (105926)",
        abs(float(sim.loc[i0, "Monto"]) - float(historico[HOJA].loc[i0, "Monto"])) < 0.5,
        f"{float(sim.loc[i0, 'Monto']):,.0f}")
    if not ok:
        raise SystemExit("⛔ ABORTA: falló alguna invariante.")
    import json
    json.dump({"s0": s0, "s1": s1, "delta_pesos": DELTA_PESOS, "dup": MONTO,
               "detalle": [[o, float(x)] for o, x in por.items()]},
              open("/private/tmp/claude-501/-Users-julianlopez/0c0e43d6-95b4-481a-a558-31041bfb4896/scratchpad/pesos/res_1444.json", "w"))

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for hh, dfh in historico.items():
            w.book.create_sheet(hh[:31])
            dfh.to_excel(w, sheet_name=hh[:31], index=False)
    buf.seek(0); data_bytes = buf.read()
    harness.clear_msgs()
    try:
        mod.guard_frescura_historico(historico)
        print("  ✅ GUARD A PASA (0 pérdidas)")
    except harness._Stop:
        for n, t in harness.MENSAJES:
            print(f"  [{n}] {t[:400]}")
        raise SystemExit("⛔ GUARD A BLOQUEÓ")

    if not ESCRIBIR:
        banner("DRY-RUN TERMINADO — 0 ESCRITURAS")
        print(f"  Para escribir: python3 correccion_1444_dup_pesos_20261021.py --escribir {md.rev}")
        return

    banner("7) SUBIDA A DROPBOX (capa C hace el respaldo)")
    md_pre = mod.dbx.files_get_metadata(cfg["remote_path"])
    if md_pre.rev != REV_ESPERADA:
        raise SystemExit(f"⛔ ABORTA SIN ESCRIBIR: el histórico se movió (rev {md_pre.rev}).")
    harness.clear_msgs()
    mod.upload_to_dropbox(data_bytes)
    backup = None
    for n, t in harness.MENSAJES:
        print(f"  [{n}] {t}")
        if "Respaldo previo creado" in t and "`" in t:
            backup = t.split("`")[1]
    harness.clear_msgs()

    banner("8) VALIDACIÓN POST-ESCRITURA")
    md2 = mod.dbx.files_get_metadata(cfg["remote_path"])
    print(f"  rev NUEVA = {md2.rev}  size = {md2.size:,}")
    print(f"  backup    = {backup}")
    _, r2 = mod.dbx.files_download(cfg["remote_path"])
    contenido = r2.content
    rel = pd.read_excel(io.BytesIO(contenido), sheet_name=None)
    ok2 = True
    def chk2(n, c, det=""):
        nonlocal ok2
        print(f"  {'✔' if c else '🚨'} {n:<64} {det}")
        ok2 = ok2 and bool(c)
    qa = rel[HOJA]["Orden"].astype(str).str.strip() == ORDEN_ORIG
    chk2("1444: cobro de duplicados -28.921.400 presente", int(qa.sum()) == 1 and abs(float(pd.to_numeric(rel[HOJA].loc[qa, "Monto"]).sum()) + MONTO) < 0.01)
    _pl = rel[HOJA].set_index(rel[HOJA]["Orden"].astype(str).str.strip())["Peso_lb"]
    chk2("1444: los 186 envíos con peso final", all(abs(float(_pl[o]) - float(p)) < 1e-9 for o, p in PF.items()))
    chk2("1444: saldo", abs(saldo(rel[HOJA]) - s1) < 0.01, f"COP {saldo(rel[HOJA]):,.2f}")
    for hoja in vivo:
        if hoja == HOJA:
            continue
        A, B = vivo[hoja], rel[hoja]
        igual = len(A) == len(B)
        if igual and "Monto" in A.columns:
            igual = abs(pd.to_numeric(A["Monto"], errors="coerce").fillna(0).sum() -
                        pd.to_numeric(B["Monto"], errors="coerce").fillna(0).sum()) < 0.01
        chk2(f"verbatim: {hoja}", igual, f"{len(A)} filas")

    banner("9) ingresos_extra.xlsx (OneDrive) — misma fila, mismo Orden")
    import openpyxl
    bk = EXTRA.replace(".xlsx", f"_backup_{pd.Timestamp.today():%Y%m%d_%H%M%S}_pre_correccion_1444.xlsx")
    shutil.copy2(EXTRA, bk)
    chk2("respaldo de ingresos_extra.xlsx", os.path.getsize(bk) == os.path.getsize(EXTRA), bk)
    wb = openpyxl.load_workbook(EXTRA); ws = wb[HEX]
    cab = [c.value for c in ws[1]]
    val = {"Tipo": "Ingreso", "Fecha": FECHA.to_pydatetime(), "Orden": ORDEN_ORIG, "Monto": -MONTO,
           "Usuario": u_orig, "Casillero": int(CAS), "Motivo": "Ingreso_extra", "Nombre del producto": NOMBRE}
    ws.append([val.get(c) for c in cab]); wb.save(EXTRA)
    ex2 = pd.read_excel(EXTRA, sheet_name=None)
    chk2("ingresos_extra.xlsx: la corrección quedó una vez", int((ex2[HEX]["Orden"].astype(str) == ORDEN_ORIG).sum()) == 1)
    chk2("ingresos_extra.xlsx: las demás hojas intactas", all(ex2[k].equals(ex[k]) for k in ex if k != HEX))

    banner("10) COPIA A ONEDRIVE")
    hoy = pd.Timestamp.today().strftime("%Y%m%d")
    destino = f"{ONEDRIVE_DIR}/{hoy}_Historico_mayoristas.xlsx"
    os.makedirs(ONEDRIVE_ANT, exist_ok=True)
    # ⚠️ COPIAR, no mover (el 18-sep OneDrive perdió un archivo movido): copiar, verificar
    # tamaño y solo entonces borrar/sobrescribir.
    for f in sorted(os.listdir(ONEDRIVE_DIR)):
        if f.endswith("_Historico_mayoristas.xlsx"):
            src_f = f"{ONEDRIVE_DIR}/{f}"
            dst_f = f"{ONEDRIVE_ANT}/{f.replace('.xlsx','')}_PRE_correccion_1444.xlsx"
            if os.path.exists(dst_f):
                raise SystemExit(f"⛔ ya existe {dst_f}; no se pisa. Revisar a mano.")
            shutil.copy2(src_f, dst_f)
            if not (os.path.exists(dst_f) and os.path.getsize(dst_f) == os.path.getsize(src_f)):
                raise SystemExit(f"⛔ la copia a Antiguos no quedó bien; NO se reemplaza {f}.")
            if f != os.path.basename(destino):
                os.remove(src_f)
            print(f"  archivada: {f} → Antiguos/{os.path.basename(dst_f)}")
    with open(destino, "wb") as fh:
        fh.write(contenido)
    print(f"  escrita:   {destino}  ({os.path.getsize(destino):,} bytes)")
    chk2("la copia de OneDrive pesa lo mismo que el vivo",
         os.path.getsize(destino) == md2.size, f"{os.path.getsize(destino):,} vs {md2.size:,}")

    print(f"\n  {'✅ CORRECCIÓN APLICADA A 1444 Y VERIFICADA' if ok2 else '🚨 REVISAR'}")
    print(f"  rev nueva: {md2.rev}\n  rollback:  {backup}")
    print("\n  ↩️  Para revertir: restaurar los backups (histórico e ingresos_extra.xlsx) y retirar PESO_FINAL_1444 de la app.")


if __name__ == "__main__":
    sys.path.insert(0, os.environ.get("HARNESS_DIR", "."))
    main()
