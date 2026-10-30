#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Cobro manual a 1444 de una compra pagada con la AAdvantage (30-oct-2026).

QEEQ.COM DUBLIN IRL · USD 243,39 · 22-ago-2026 · cargo #1124 de Iris (movimientos_tc).
La AAdvantage no tiene módulo (no es tarjeta de mayoristas), así que el cargo nunca entró al
histórico. Julián lo cerró en Iris como compra mayorista de 1444 y pidió subirlo aquí.

Se carga como Ingreso_extra NEGATIVO en su fecha real (mismo patrón que los cobros Correal),
con Orden propio y también en ingresos_extra.xlsx. COP = USD × TRM de _amex_trm_dia (oficial
+125) del 22-ago. No genera incentivo: el Motivo no es de tarjeta.
La comisión quincenal NO se toca (las quincenas desde 16-31 ago están cerradas); el dry-run
informa cuánto cambiaría si se recalculara.

  sin argumentos  -> dry-run (0 escrituras)
  --escribir REV  -> escribe a Dropbox (REV = rev del dry-run) y refresca la copia de OneDrive
"""
import sys, os, io, shutil, hashlib, warnings
warnings.filterwarnings("ignore")
import pandas as pd

ESCRIBIR = "--escribir" in sys.argv
REV_ESPERADA = sys.argv[sys.argv.index("--escribir") + 1] if ESCRIBIR else None
CAS = "1444"
FECHA = "2026-08-22"
USD = 243.39
ORDEN = "compra_aadvantage_qeeq_20260822_1444"
EXTRA = "/Users/julianlopez/Library/CloudStorage/OneDrive-Personal/Historico Carga/Conciliacion/Mayoristas/ingresos_extra.xlsx"
ONEDRIVE_DIR = "/Users/julianlopez/Library/CloudStorage/OneDrive-Personal/Historico Carga/Conciliacion/Mayoristas"
ONEDRIVE_ANT = f"{ONEDRIVE_DIR}/Antiguos"


def saldo(d):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    return float(pd.to_numeric(t["Monto"], errors="coerce").iloc[-1]) if len(t) else float("nan")


def usuario_de_totales(d):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    u = t["Usuario"].astype(str).str.strip()
    u = u[~u.str.lower().isin({"", "nan", "none"})]
    return u.mode().iloc[0] if len(u) else ""


def cas_de_totales(d, fb):
    t = d[d["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"]
    v = t["Casillero"].dropna()
    return v.iloc[-1] if len(v) else fb


def tots(df):
    t = df[df["Tipo"].astype(str).str.strip().str.upper() == "TOTAL"].copy()
    return t.assign(F=pd.to_datetime(t["Fecha"]).dt.normalize()).groupby("F")["Monto"].last().astype(float)


def main():
    import harness
    mod = harness.cargar_app()
    SEP = "=" * 94
    def banner(t): print(f"\n{SEP}\n{t}\n{SEP}")
    ok = True
    def chk(n, c, det=""):
        nonlocal ok
        print(f"  {'✔' if c else '🚨'} {n:<70} {det}")
        ok = ok and bool(c)
    print(f"MODO: {'🔴 ESCRITURA REAL' if ESCRIBIR else '🟢 DRY-RUN (0 escrituras)'}")
    for fn in ("procesar_egresos", "procesar_amex", "procesar_rakuten", "procesar_robinhood",
               "procesar_capital", "procesar_usbank", "procesar_intuit", "procesar_discover",
               "agregar_incentivo_amex"):
        def _boom(*a, _n=fn, **k):
            raise AssertionError(f"{_n} fue llamada — este script solo agrega un cobro")
        setattr(mod, fn, _boom)

    trm = mod._amex_trm_dia(FECHA, {})
    chk(f"TRM del {FECHA} (oficial +125)", trm is not None and trm > 0, f"{trm}")
    if not ok:
        raise SystemExit("⛔ ABORTA: sin TRM.")
    monto = -round(USD * trm, 2)
    nombre = f"Compra QEEQ.COM (Dublin) USD {USD:,.2f} con AAdvantage — cargo Iris #1124"
    print(f"  COP = {USD} × {trm:,.2f} = {-monto:,.2f}")

    banner("1) HISTÓRICO VIVO FRESCO")
    cfg = mod.st.secrets["dropbox"]
    md = mod.dbx.files_get_metadata(cfg["remote_path"])
    _, res = mod.dbx.files_download(cfg["remote_path"])
    vivo = pd.read_excel(io.BytesIO(res.content), sheet_name=None)
    HOJA = next(h for h in vivo if h.split(" - ")[0].strip() == CAS and h.strip().endswith("Moises"))
    s0 = saldo(vivo[HOJA])
    print(f"  rev={md.rev}  modificado={md.server_modified}  size={md.size:,}")
    print(f"  '{HOJA}': saldo {s0:,.2f}")
    if ESCRIBIR and md.rev != REV_ESPERADA:
        raise SystemExit(f"⛔ ABORTA: el histórico se movió (esperaba {REV_ESPERADA}). Rehacer dry-run.")

    banner("2) VERIFICACIONES PREVIAS")
    # En 1444 hay compras/reembolsos Qeeq con la ROBINHOOD del 17 al 20-sep (otras compras); por
    # eso la búsqueda se limita a 15-ago..10-sep, que es donde podría estar ESTE cargo.
    for h, dd in vivo.items():
        if "Orden" not in dd.columns:
            continue
        fq = pd.to_datetime(dd["Fecha"], errors="coerce").between("2026-08-15", "2026-09-10")
        txt = dd.astype(str).apply(lambda c: c.str.contains("QEEQ", case=False, na=False)).any(axis=1)
        chk(f"'{h}' no tiene QEEQ entre 15-ago y 10-sep", not (txt & fq).any(), f"{int((txt & fq).sum())}")
    d0 = vivo[HOJA]
    f0 = pd.to_datetime(d0["Fecha"], errors="coerce")
    usd0 = pd.to_numeric(d0["Monto"], errors="coerce") / pd.to_numeric(d0["TRM"], errors="coerce")
    tipo0 = d0["Tipo"].astype(str).str.strip().str.upper()
    cobro0 = (tipo0 == "EGRESO") | ((d0["Motivo"].astype(str).str.strip() == "Ingreso_extra")
                                    & (pd.to_numeric(d0["Monto"], errors="coerce") < 0))
    cerca = cobro0 & f0.between("2026-08-15", "2026-09-10") & ((usd0.abs() - USD).abs() < 1.5)
    chk("1444 no tiene otro COBRO de ~USD 243,39 entre 15-ago y 10-sep", not cerca.any(), f"{int(cerca.sum())}")
    ex = pd.read_excel(EXTRA, sheet_name=None)
    HEX = next(h for h in ex if h.split("-")[0].strip() == CAS)
    chk("ingresos_extra.xlsx todavía no lo tiene", not ex[HEX]["Orden"].astype(str).str.strip().eq(ORDEN).any())
    if not ok:
        raise SystemExit("⛔ ABORTA.")

    banner(f"3) APLICAR — {FECHA}  {ORDEN}")
    hoy = pd.Timestamp.today().strftime("%Y-%m-%d")
    historico = {k: mod.asegurar_columnas_historico(v.copy()) for k, v in vivo.items()}
    d = historico[HOJA]
    u_orig = usuario_de_totales(vivo[HOJA]) or "Maria Moises"
    cas_t = cas_de_totales(vivo[HOJA], CAS)
    fila = {c: "" for c in d.columns}
    fila.update({"Fecha": pd.Timestamp(FECHA), "Tipo": "Ingreso", "Monto": monto, "Orden": ORDEN,
                 "Motivo": "Ingreso_extra", "TRM": trm, "Usuario": u_orig, "Casillero": int(CAS),
                 "Nombre del producto": nombre, "Fecha de Carga": hoy})
    print(f"  + {FECHA} {ORDEN}  {monto:,.2f}  {nombre}")
    d = pd.concat([d, pd.DataFrame([fila])], ignore_index=True)
    historico[HOJA] = mod.recalcular_totales_diarios(d, usuario=u_orig, cas=cas_t)

    banner("3b) COMISIÓN — NO se toca; cuánto cambiaría si se recalculara")
    ta, tb = tots(vivo[HOJA]), tots(historico[HOJA])
    QS = [("16-fin agosto 2026", "2026-08-16", "2026-08-31"), ("1-15 septiembre 2026", "2026-09-01", "2026-09-15"),
          ("16-fin septiembre 2026", "2026-09-16", "2026-09-30"), ("1-15 octubre 2026", "2026-10-01", "2026-10-15")]
    dcom = 0.0
    for et, a, b in QS:
        pa = ta[(ta.index >= a) & (ta.index <= b)].min(); pb = tb[(tb.index >= a) & (tb.index <= b)].min()
        ca = round(abs(pa) * 0.015, 2) if pa < 0 else 0.0; cb = round(abs(pb) * 0.015, 2) if pb < 0 else 0.0
        dcom += cb - ca
        print(f"  {et:<24} peor {pa:>16,.2f} → {pb:>16,.2f}   1,5%: {ca:>12,.2f} → {cb:>12,.2f}  (Δ {cb - ca:+,.2f})")
    print(f"  Δ comisión si se recalculara: {dcom:+,.2f}  (NO se aplica)")

    banner("4) INVARIANTES")
    s1 = saldo(historico[HOJA])
    chk("el saldo baja exactamente el cobro", abs((s0 - s1) + monto) < 0.01, f"{s0:,.2f} → {s1:,.2f}")
    def movs(df):
        return int((df["Tipo"].astype(str).str.strip().str.upper() != "TOTAL").sum())
    chk("+1 fila, no se borra nada", movs(historico[HOJA]) == movs(vivo[HOJA]) + 1)
    antes = ta.index[ta.index < pd.Timestamp(FECHA)]
    chk(f"saldos diarios anteriores al {FECHA} idénticos", (tb.reindex(antes) - ta.reindex(antes)).abs().max() < 0.01, f"{len(antes)} días")
    def comis(df):
        c = df[df["Orden"].astype(str).str.contains("Comision", case=False, na=False)]
        return sorted(zip(c["Orden"].astype(str), pd.to_numeric(c["Monto"]).round(2)))
    chk("comisiones idénticas", comis(vivo[HOJA]) == comis(historico[HOJA]), f"{len(comis(vivo[HOJA]))}")
    vac = {"", "nan", "none", "nat"}
    a = {y for y in vivo[HOJA]["Orden"].astype(str).str.strip() if y.lower() not in vac}
    b = {y for y in historico[HOJA]["Orden"].astype(str).str.strip() if y.lower() not in vac}
    chk(f"{HOJA}: 0 Orden perdidos", not (a - b), f"{len(a - b)}")
    for hoja in vivo:
        if hoja == HOJA:
            continue
        A, B = vivo[hoja], historico[hoja]
        igual = len(A) == len(B)
        if igual and "Monto" in A.columns:
            igual = abs(pd.to_numeric(A["Monto"], errors="coerce").fillna(0).sum() -
                        pd.to_numeric(B["Monto"], errors="coerce").fillna(0).sum()) < 0.01
        chk(f"verbatim: {hoja}", igual, f"{len(A)} filas")
    if not ok:
        raise SystemExit("⛔ ABORTA: falló alguna invariante.")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for hh, dfh in historico.items():
            w.book.create_sheet(hh[:31])
            dfh.to_excel(w, sheet_name=hh[:31], index=False)
    buf.seek(0); data_bytes = buf.read()
    harness.clear_msgs()
    try:
        mod.guard_frescura_historico(historico)
        print("  ✅ GUARD A PASA (0 pérdidas)")
    except harness._Stop:
        for n, t in harness.MENSAJES:
            print(f"  [{n}] {t[:400]}")
        raise SystemExit("⛔ GUARD A BLOQUEÓ")

    if not ESCRIBIR:
        banner("DRY-RUN TERMINADO — 0 ESCRITURAS")
        print(f"  Para escribir: python3 cobro_qeeq_aadvantage_1444_20261030.py --escribir {md.rev}")
        return

    banner("5) SUBIDA A DROPBOX (capa C hace el respaldo)")
    md_pre = mod.dbx.files_get_metadata(cfg["remote_path"])
    if md_pre.rev != REV_ESPERADA:
        raise SystemExit(f"⛔ ABORTA SIN ESCRIBIR: el histórico se movió (rev {md_pre.rev}).")
    harness.clear_msgs()
    mod.upload_to_dropbox(data_bytes)
    backup = None
    for n, t in harness.MENSAJES:
        print(f"  [{n}] {t}")
        if "Respaldo previo creado" in t and "`" in t:
            backup = t.split("`")[1]
    harness.clear_msgs()

    banner("6) VALIDACIÓN POST-ESCRITURA")
    md2 = mod.dbx.files_get_metadata(cfg["remote_path"])
    print(f"  rev NUEVA = {md2.rev}  size = {md2.size:,}")
    print(f"  backup    = {backup}")
    _, r2 = mod.dbx.files_download(cfg["remote_path"])
    contenido = r2.content
    rel = pd.read_excel(io.BytesIO(contenido), sheet_name=None)
    ok2 = True
    def chk2(n, c, det=""):
        nonlocal ok2
        print(f"  {'✔' if c else '🚨'} {n:<64} {det}")
        ok2 = ok2 and bool(c)
    q = rel[HOJA]["Orden"].astype(str).str.strip() == ORDEN
    chk2(f"1444: {ORDEN}", int(q.sum()) == 1 and abs(float(pd.to_numeric(rel[HOJA].loc[q, "Monto"]).iloc[0]) - monto) < 0.01)
    chk2("1444: saldo", abs(saldo(rel[HOJA]) - s1) < 0.01, f"COP {saldo(rel[HOJA]):,.2f}")
    for hoja in vivo:
        if hoja == HOJA:
            continue
        A, B = vivo[hoja], rel[hoja]
        igual = len(A) == len(B)
        if igual and "Monto" in A.columns:
            igual = abs(pd.to_numeric(A["Monto"], errors="coerce").fillna(0).sum() -
                        pd.to_numeric(B["Monto"], errors="coerce").fillna(0).sum()) < 0.01
        chk2(f"verbatim: {hoja}", igual, f"{len(A)} filas")

    banner("7) ingresos_extra.xlsx (OneDrive)")
    import openpyxl
    bk = EXTRA.replace(".xlsx", f"_backup_{pd.Timestamp.today():%Y%m%d_%H%M%S}_pre_qeeq.xlsx")
    shutil.copy2(EXTRA, bk)
    chk2("respaldo de ingresos_extra.xlsx", os.path.getsize(bk) == os.path.getsize(EXTRA), bk)
    wb = openpyxl.load_workbook(EXTRA); ws = wb[HEX]
    cab = [c.value for c in ws[1]]
    val = {"Tipo": "Ingreso", "Fecha": pd.Timestamp(FECHA).to_pydatetime(), "Orden": ORDEN, "Monto": monto,
           "Usuario": u_orig, "Casillero": int(CAS), "Motivo": "Ingreso_extra", "Nombre del producto": nombre}
    ws.append([val.get(c) for c in cab])
    wb.save(EXTRA)
    ex2 = pd.read_excel(EXTRA, sheet_name=None)
    chk2("ingresos_extra.xlsx: quedó una vez", int((ex2[HEX]["Orden"].astype(str) == ORDEN).sum()) == 1)
    chk2("ingresos_extra.xlsx: las demás hojas intactas", all(ex2[k].equals(ex[k]) for k in ex if k != HEX))

    banner("8) COPIA A ONEDRIVE")
    hoyc = pd.Timestamp.today().strftime("%Y%m%d")
    destino = f"{ONEDRIVE_DIR}/{hoyc}_Historico_mayoristas.xlsx"
    os.makedirs(ONEDRIVE_ANT, exist_ok=True)
    # ⚠️ COPIAR, no mover (el 18-sep OneDrive perdió un archivo movido).
    for f in sorted(os.listdir(ONEDRIVE_DIR)):
        if f.endswith("_Historico_mayoristas.xlsx"):
            src_f = f"{ONEDRIVE_DIR}/{f}"
            dst_f = f"{ONEDRIVE_ANT}/{f.replace('.xlsx','')}_PRE_qeeq.xlsx"
            if os.path.exists(dst_f):
                raise SystemExit(f"⛔ ya existe {dst_f}; no se pisa. Revisar a mano.")
            shutil.copy2(src_f, dst_f)
            if not (os.path.exists(dst_f) and os.path.getsize(dst_f) == os.path.getsize(src_f)):
                raise SystemExit(f"⛔ la copia a Antiguos no quedó bien; NO se reemplaza {f}.")
            if f != os.path.basename(destino):
                os.remove(src_f)
            print(f"  archivada: {f} → Antiguos/{os.path.basename(dst_f)}")
    with open(destino, "wb") as fh:
        fh.write(contenido)
    print(f"  escrita:   {destino}  ({os.path.getsize(destino):,} bytes)")
    chk2("la copia de OneDrive es idéntica al vivo (sha256)",
         hashlib.sha256(open(destino, "rb").read()).hexdigest() == hashlib.sha256(contenido).hexdigest())

    print(f"\n  {'✅ COBRO QEEQ APLICADO A 1444 Y VERIFICADO' if ok2 else '🚨 REVISAR'}")
    print(f"  rev nueva: {md2.rev}\n  rollback:  {backup}")


if __name__ == "__main__":
    sys.path.insert(0, os.environ.get("HARNESS_DIR", "."))
    main()
